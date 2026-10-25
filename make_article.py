"""Build the illustrated reading copy and editable Word article from audited data."""
import base64
import html
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
STEM='Learning-in-a-Separate-Neural-Memory'
ARTICLE=json.loads((ROOT/'article_content.json').read_text())
RESULT=json.loads((ROOT/'results/summary.json').read_text())
NAMES={'clean':'clean_memory','untrained':'untrained_memory','empty':'no_write',
       'poison':'after_false_updates','correction':'after_verified_corrections'}
VALUES={k:f"{RESULT['conditions'][v]['accuracy']*100:.1f}" for k,v in NAMES.items()}
VALUES['capacity64']=f"{RESULT['capacity']['64']*100:.1f}"

class RichText(HTMLParser):
    def __init__(self,paragraph):
        super().__init__();self.p=paragraph;self.bold=0
    def handle_starttag(self,tag,attrs):
        if tag=='strong':self.bold+=1
    def handle_endtag(self,tag):
        if tag=='strong':self.bold-=1
    def handle_data(self,text):
        self.p.add_run(text).bold=bool(self.bold)

def add_link(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'087F79');props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'17');props.append(size)
    run.append(props);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)

def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(8)
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles['Title'].font.name='Calibri';doc.styles['Title'].font.size=Pt(26)
    doc.styles['Title'].paragraph_format.space_after=Pt(9)
    doc.styles['Subtitle'].font.size=Pt(12)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(9)
    doc.styles['Caption'].font.name='Calibri';doc.styles['Caption'].font.size=Pt(8.5)
    doc.styles['Caption'].font.color.rgb=RGBColor.from_string('536773')
    doc.styles['Caption'].paragraph_format.space_after=Pt(12)
    doc.core_properties.author=ARTICLE['author'];doc.core_properties.title=ARTICLE['title']
    doc.add_paragraph(ARTICLE['title'],'Title')
    doc.add_paragraph(ARTICLE['subtitle'],'Subtitle')
    p=doc.add_paragraph(ARTICLE['author']+' · partha_systems.ai');p.runs[0].font.size=Pt(9)
    content=[f'<header><p class="eyebrow">TITANS INSPIRED MEMORY · AN EXPERIMENT</p><h1>{html.escape(ARTICLE["title"])}</h1><p class="subtitle">{html.escape(ARTICLE["subtitle"])}</p><p class="byline">{html.escape(ARTICLE["author"])} · partha_systems.ai</p></header>']
    md=['# '+ARTICLE['title'],'',ARTICLE['subtitle'],'','**'+ARTICLE['author']+'**','']
    image_count=0;texts=[]
    for block in ARTICLE['blocks']:
        kind=block['type']
        if kind=='break':doc.add_page_break();continue
        if kind in ['p','note']:
            text=block['text'].format(**VALUES);texts.append(re.sub('<[^>]+>','',text))
            p=doc.add_paragraph();RichText(p).feed(text)
            if kind=='note':
                for run in p.runs:run.font.size=Pt(8.5)
            content.append('<p'+(' class="note"' if kind=='note' else '')+'>'+text+'</p>')
            md.extend([text.replace('<strong>','**').replace('</strong>','**'),''])
        elif kind=='image':
            image_count+=1;path=ROOT/'figures'/block['file'];data=path.read_bytes()
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(path),width=Inches(6.85))
            shape._inline.docPr.set('descr',block['alt'])
            doc.add_paragraph(f'Figure {image_count}. '+block['caption'],'Caption')
            src='data:image/png;base64,'+base64.b64encode(data).decode()
            content.append(f'<figure><img src="{src}" alt="{html.escape(block["alt"],quote=True)}"><figcaption>Figure {image_count}. {html.escape(block["caption"])}</figcaption></figure>')
            md.extend([f'![{block["alt"]}](figures/{block["file"]})','',f'*Figure {image_count}. {block["caption"]}*',''])
    content.append('<footer><p>Research sources</p><ol>')
    md.extend(['## Research sources',''])
    for i,source in enumerate(ARTICLE['sources'],1):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        add_link(p,f'[{i}] '+source['label'],source['url'])
        content.append(f'<li><a href="{source["url"]}">{html.escape(source["label"])}</a></li>')
        md.append(f'{i}. [{source["label"]}]({source["url"]})')
    content.append('</ol></footer>')
    doc.save(ROOT/(STEM+'.docx'))
    css='''*{box-sizing:border-box}body{margin:0;background:#f8f7f3;color:#17313d;font-family:Georgia,serif}article{max-width:1080px;margin:36px auto;padding:44px;background:white;border:1px solid #dfe4de;border-radius:10px}header,p,footer{max-width:735px;margin-left:auto;margin-right:auto}h1{font-size:46px;line-height:1.12;letter-spacing:-1px;color:#111;margin:14px 0}.eyebrow,.byline,figcaption,footer{font-family:Arial,sans-serif}.eyebrow{font-size:12px;letter-spacing:1.5px;color:#087f79;font-weight:700}.subtitle{font-size:22px;color:#536773;line-height:1.4}.byline{font-size:13px;color:#536773;margin-bottom:30px}p{font-size:19px;line-height:1.7;margin-top:0;margin-bottom:20px}figure{margin:30px 0 34px}img{display:block;width:100%;height:auto;border:1px solid #dfe4de;border-radius:5px}figcaption{max-width:880px;margin:10px auto 0;font-size:13px;line-height:1.5;color:#536773}.note{font-size:14px;color:#536773}a{color:#087f79}footer{margin-top:24px;font-size:13px;line-height:1.6}footer p{font-size:14px;font-weight:700}footer ol{padding-left:20px}@media(max-width:680px){article{margin:0;padding:25px 18px;border:0;border-radius:0}h1{font-size:34px}.subtitle{font-size:20px}p{font-size:18px}figure{margin:26px -8px}figcaption{padding:0 8px}}'''
    document='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(ARTICLE['title'])+'</title><style>'+css+'</style></head><body><article>'+''.join(content)+'</article></body></html>'
    (ROOT/(STEM+'.html')).write_text(document)
    (ROOT/'article.md').write_text('\n'.join(md)+'\n')
    assert document.count('data:image/png;base64,')==image_count
    print(json.dumps(dict(article_words_without_captions=len(' '.join(texts).split()),images=image_count,docx=str(ROOT/(STEM+'.docx')),html=str(ROOT/(STEM+'.html')))))

if __name__=='__main__':main()
